"""Module chuẩn hóa đề thi tự động từ file Word thô."""
import os
import re
import logging
from typing import Dict, List, Optional, Tuple
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell

logger = logging.getLogger(__name__)


def iter_block_items(parent):
    if parent.__class__.__name__ == 'Document':
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent type")

    for child in parent_elm.iterchildren():
        if child.tag.endswith('p'):
            yield Paragraph(child, parent)
        elif child.tag.endswith('tbl'):
            yield Table(child, parent)


def get_all_paragraphs(doc) -> List[Paragraph]:
    paragraphs = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            paragraphs.append(block)
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for nested_block in iter_block_items(cell):
                        if isinstance(nested_block, Paragraph):
                            paragraphs.append(nested_block)
    return paragraphs


def split_paragraphs_at_br(doc):
    """Tách triệt để một đoạn văn thành nhiều đoạn văn nếu nó chứa ngắt dòng mềm (Soft Return / w:br)."""
    import copy
    from docx.oxml import OxmlElement
    for p in doc.paragraphs:
        p_elem = p._element
        has_br = False
        for r in p_elem.findall(qn('w:r')):
            if r.find(qn('w:br')) is not None:
                has_br = True
                break
                
        if not has_br:
            continue
            
        parent = p_elem.getparent()
        if parent is None:
            continue
            
        current_p = OxmlElement('w:p')
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is not None:
            current_p.append(copy.deepcopy(pPr))
            
        new_paras = [current_p]
        
        for child in list(p_elem):
            if child.tag == qn('w:pPr'):
                continue
            elif child.tag == qn('w:r'):
                brs = child.findall(qn('w:br'))
                if not brs:
                    current_p.append(copy.deepcopy(child))
                else:
                    current_r = OxmlElement('w:r')
                    rPr = child.find(qn('w:rPr'))
                    if rPr is not None:
                        current_r.append(copy.deepcopy(rPr))
                        
                    current_p.append(current_r)
                    
                    for r_child in list(child):
                        if r_child.tag == qn('w:rPr'):
                            continue
                        elif r_child.tag == qn('w:br'):
                            current_p = OxmlElement('w:p')
                            if pPr is not None:
                                current_p.append(copy.deepcopy(pPr))
                            new_paras.append(current_p)
                            
                            current_r = OxmlElement('w:r')
                            if rPr is not None:
                                current_r.append(copy.deepcopy(rPr))
                            current_p.append(current_r)
                        else:
                            current_r.append(copy.deepcopy(r_child))
            else:
                current_p.append(copy.deepcopy(child))
                
        for np in new_paras:
            p_elem.addprevious(np)
            
        parent.remove(p_elem)


def convert_choice_tables_to_paragraphs(doc):
    """Chuyển đổi các bảng chứa các phương án trắc nghiệm thành các đoạn văn bình thường trong thân văn bản."""
    CHOICE_LABEL_RE = re.compile(r'^\s*(?:\()?([A-Ha-h])\s*[.:)/\-]\s*')
    
    tables_to_process = []
    for table in doc.tables:
        choice_cells_count = 0
        total_cells = 0
        for row in table.rows:
            for cell in row.cells:
                total_cells += 1
                text = cell.text.strip()
                if CHOICE_LABEL_RE.match(text):
                    choice_cells_count += 1
                    
        # Nếu bảng có trên 50% số ô bắt đầu bằng nhãn lựa chọn và có ít nhất 2 ô khớp
        if choice_cells_count >= 2 and choice_cells_count >= total_cells * 0.5:
            tables_to_process.append(table)
            
    for table in tables_to_process:
        tbl_elem = table._element
        parent = tbl_elem.getparent()
        if parent is None:
            continue
            
        paragraphs_to_move = []
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    paragraphs_to_move.append(p)
                    
        for p in paragraphs_to_move:
            p_elem = p._element
            tbl_elem.addprevious(p_elem)
            
        parent.remove(tbl_elem)


def clean_empty_paragraphs(doc: Document):
    """Xóa bỏ triệt để các đoạn văn rỗng rác (không text, không ảnh, không math)."""
    for p in list(doc.paragraphs):
        if not p.text.strip():
            p_elem = p._element
            if p_elem.findall('.//' + qn('w:drawing')):
                continue
            if p_elem.findall('.//' + qn('w:pict')):
                continue
            if p_elem.findall('.//' + qn('m:oMath')):
                continue
            if p_elem.findall('.//' + qn('w:object')):
                continue
            parent = p_elem.getparent()
            if parent is not None:
                parent.remove(p_elem)


# Các biểu thức chính quy (Regex) để nhận dạng
# Nhận dạng Phần/Part: Phần I, Phần 1, Part 1, 2. Chọn Đ (Đúng), 3. Trắc nghiệm trả lời ngắn
QUESTION_RE = re.compile(r'^(Câu|Question|Bài|câu|bài)\s*(hỏi\s+)?(\d+)\s*(?:\([^)]*\))?\s*[:.\-)?\s]?', re.IGNORECASE)
PLAIN_NUMBER_RE = re.compile(r'^(\d+)\s*(?:\([^)]*\))?\s*[.:)]\s*(?!\d)')
SECTION_RE = re.compile(
    r'^\s*(?:'
    r'(phần|part)\s*([IVX]+|\d+)'
    r'|'
    r'([IVX]+|\d+)\s*[.:]\s*(?=.*?(?:trắc nghiệm|tự luận|đúng[\s/]*sai|trả lời ngắn|chọn đ\s*\())'
    r')', re.IGNORECASE
)

INLINE_MC_RE = re.compile(r'(?<![a-zA-Z])(?:\()?([A-H])\s*[.:)/\-]\s*')
INLINE_TF_RE = re.compile(r'(?<![a-zA-Z])(?:\()?([a-h])\s*[.:)/\-]\s*')


def is_short_answer_value(val: str) -> bool:
    """Kiểm tra xem giá trị đáp án có phải là đáp án tự luận ngắn hay không (không phải A/B/C/D và không phải Đúng/Sai)."""
    if not val:
        return False
    val_clean = val.strip().upper()
    # Nếu là phương án trắc nghiệm đơn (A, B, C, D...)
    if len(val_clean) == 1 and 'A' <= val_clean <= 'H':
        return False
    # Nếu là chuỗi Đúng/Sai (Đ,S,Đ,S...)
    if re.match(r'^[ĐSDTRUEFSAI](?:[\s,]+[ĐSDTRUEFSAI])*$', val_clean):
        return False
    return True

def _parse_table_answer_key(table) -> Dict[int, str]:
    """Trích xuất đáp án từ bảng ở cuối đề, hỗ trợ cả bảng dọc và bảng ngang (Phần 1, 2, 3)."""
    keys = {}
    
    # Lấy toàn bộ văn bản trong bảng dưới dạng lưới (grid) của các dòng và cột
    grid = []
    for row in table.rows:
        grid.append([cell.text.strip() for cell in row.cells])
        
    if not grid or not grid[0]:
        return keys
        
    num_rows = len(grid)
    num_cols = len(grid[0])
    
    # 1. Kiểm tra xem có dòng nào chứa số câu hỏi (Question Row) không
    # Một dòng câu hỏi hợp lệ là dòng có đa số các ô chứa số (trừ ô đầu tiên thường là nhãn "Câu")
    question_rows = []
    for r_idx in range(num_rows):
        row_cells = grid[r_idx]
        if not row_cells:
            continue
        is_col0_q = re.match(r'^(?:Câu\s*)?(\d+)$', row_cells[0], re.IGNORECASE) is not None
        start_idx = 0 if is_col0_q else 1
        
        digits_count = sum(1 for c in row_cells[start_idx:] if re.match(r'^(?:Câu\s*)?(\d+)$', c, re.IGNORECASE))
        non_empty_cols = sum(1 for c in row_cells[start_idx:] if c)
        if digits_count >= 2 and digits_count >= non_empty_cols * 0.7:
            question_rows.append(r_idx)
            
    if question_rows:
        # Bảng ngang!
        for q_row_idx in question_rows:
            # Xác định các số câu hỏi trong dòng này
            q_numbers = []
            col_mapping = {} # col_idx -> q_num
            
            start_col = 0 if grid[q_row_idx] and re.match(r'^(?:Câu\s*)?(\d+)$', grid[q_row_idx][0], re.IGNORECASE) else 1
            for c_idx in range(start_col, num_cols):
                val = grid[q_row_idx][c_idx]
                m = re.match(r'^(?:Câu\s*)?(\d+)$', val, re.IGNORECASE)
                if m:
                    q_num = int(m.group(1))
                    q_numbers.append(q_num)
                    col_mapping[c_idx] = q_num
            
            # Bây giờ quét các dòng bên dưới dòng câu hỏi này
            # Check xem có các dòng chứa nhãn a), b), c), d) bên dưới không
            tf_rows = []
            for r_idx in range(q_row_idx + 1, num_rows):
                if r_idx in question_rows:
                    break
                found_label = None
                for cell_txt in grid[r_idx]:
                    cleaned = cell_txt.strip().lower()
                    if re.match(r'^[a-h]\s*[\).:]', cleaned):
                        found_label = cleaned[0]
                        break
                if found_label:
                    tf_rows.append((found_label, r_idx))
                    
            if len(tf_rows) >= 2:
                # Đây là bảng Đúng/Sai Phần 2!
                # Gom đáp án theo từng cột
                # Sắp xếp tf_rows theo nhãn a, b, c, d
                tf_rows.sort(key=lambda x: x[0])
                for c_idx, q_num in col_mapping.items():
                    ans_list = []
                    for label, r_idx in tf_rows:
                        if c_idx < len(grid[r_idx]):
                            val = grid[r_idx][c_idx].strip().upper()
                            if val in ['Đ', 'ĐÚNG', 'T', 'TRUE', 'D']:
                                ans_list.append('Đ')
                            elif val in ['S', 'SAI', 'F', 'FALSE']:
                                ans_list.append('S')
                            else:
                                ans_list.append(val)
                    if ans_list:
                        keys[q_num] = ",".join(ans_list)
            else:
                # Đây là bảng Trắc nghiệm hoặc Trả lời ngắn nằm ngang (chỉ có 1 dòng đáp án bên dưới)
                for r_idx in range(q_row_idx + 1, num_rows):
                    if r_idx in question_rows:
                        break
                    start_col = 0 if grid[q_row_idx] and re.match(r'^(?:Câu\s*)?(\d+)$', grid[q_row_idx][0], re.IGNORECASE) else 1
                    if not any(grid[r_idx][start_col:]):
                        continue
                    for c_idx, q_num in col_mapping.items():
                        if c_idx < len(grid[r_idx]):
                            ans_val = grid[r_idx][c_idx].strip()
                            if ans_val:
                                keys[q_num] = ans_val.upper()
                    break
        return keys

    # 2. Kiểm tra bảng dọc Đúng/Sai (Cột 0 là số câu, các cột sau chứa a-S, b-Đ...)
    for row in table.rows:
        cells = row.cells
        if len(cells) >= 2:
            t_first = cells[0].text.strip()
            if t_first.isdigit():
                q_num = int(t_first)
                tf_vals = []
                is_tf = False
                for c_idx in range(1, len(cells)):
                    cell_txt = cells[c_idx].text.strip()
                    m = re.match(r'^([a-h])\s*[\-.:\s]\s*([ĐSđsD])$', cell_txt, re.IGNORECASE)
                    if m:
                        ans_val = 'Đ' if m.group(2).upper() in ['Đ', 'D'] else 'S'
                        opt_idx = ord(m.group(1).lower()) - ord('a')
                        tf_vals.append((opt_idx, ans_val))
                        is_tf = True
                    else:
                        val = cell_txt.upper()
                        if val in ['Đ', 'ĐÚNG', 'T', 'TRUE', 'D']:
                            tf_vals.append((c_idx - 1, 'Đ'))
                        elif val in ['S', 'SAI', 'F', 'FALSE']:
                            tf_vals.append((c_idx - 1, 'S'))
                if is_tf and tf_vals:
                    tf_vals.sort(key=lambda x: x[0])
                    keys[q_num] = ",".join(x[1] for x in tf_vals)
                    continue

    # 3. Nếu không có dòng câu hỏi ngang, thử quét kiểu bảng dọc (2 cột: Cột lẻ là Số câu, Cột chẵn kế tiếp là Đáp án)
    for row in table.rows:
        cells = row.cells
        for i in range(len(cells) - 1):
            t_left = cells[i].text.strip()
            t_right = cells[i+1].text.strip()
            if t_left.isdigit():
                q_num = int(t_left)
                if q_num not in keys:
                    if t_right and len(t_right) < 15:
                        keys[q_num] = t_right.upper()
                    
    # 4. Quét kiểu gộp ô "1. A" hoặc "Câu 1: B" hoặc "Câu 1: 1,5"
    pat_cell = re.compile(r'^(?:Câu\s*)?(\d+)\s*[-.:=/\s]+\s*(.+)$', re.IGNORECASE)
    for row in table.rows:
        for cell in row.cells:
            txt = cell.text.strip()
            if re.match(r'^-?\d+[\.,]\d+$', txt):
                continue
            m = pat_cell.match(txt)
            if m:
                q_num = int(m.group(1))
                if q_num not in keys:
                    ans_val = m.group(2).strip()
                    keys[q_num] = ans_val
                
    return keys


def _parse_text_answer_key(text: str) -> Dict[int, str]:
    """Trích xuất đáp án từ một chuỗi văn bản."""
    keys = {}
    
    # 1. Match Part 2 (Đúng/Sai): VD: 23(Đ,S,Đ,S) hoặc 23(Đ-S-Đ-S)
    pattern_tf = re.compile(r'(\d+)\s*[\(\[:]?\s*([ĐSđsD]+(?:[\s,.\-/]+[ĐSđsD]+)+)\s*[\)\]]?', re.IGNORECASE)
    for m in pattern_tf.finditer(text):
        q_num = int(m.group(1))
        raw_vals = re.findall(r'[ĐSđsD]', m.group(2).upper())
        raw_vals = ['Đ' if v == 'D' else v for v in raw_vals]
        keys[q_num] = ",".join(raw_vals)
        
    # 2. Match Part 3 (Trả lời ngắn): VD: 26: 15.5 hoặc 26 = 15.5
    pattern_short = re.compile(r'(\d+)\s*[:=]\s*(-?\d+(?:[\.,]\d+)?)')
    for m in pattern_short.finditer(text):
        q_num = int(m.group(1))
        if q_num not in keys:
            keys[q_num] = m.group(2)
            
    # 3. Match Part 1 (Trắc nghiệm A/B/C/D): VD: 1A, 2.B, 3-C, Câu 4: D
    pattern_mc = re.compile(r'(?:Câu\s*)?(\d+)\s*[.\-:]?\s*([A-D])(?![a-zA-Z0-9])', re.IGNORECASE)
    for m in pattern_mc.finditer(text):
        q_num = int(m.group(1))
        if q_num not in keys:
            keys[q_num] = m.group(2).upper()
            
    return keys


def is_answer_title(text: str) -> bool:
    text_clean = text.strip()
    if not text_clean:
        return False
    is_title = (
        re.search(r'^\s*(?:BẢNG\s+)?ĐÁP\s+ÁN(?:\s+CHI\s+TIẾT)?\s*$', text_clean, re.IGNORECASE) or 
        re.search(r'^\s*HƯỚNG\s+DẪN\s+CHẤM(?:\s+VÀ\s+BIỂU\s+ĐIỂM)?\s*$', text_clean, re.IGNORECASE) or
        re.search(r'^\s*PHẦN\s+ĐÁP\s+ÁN\s*$', text_clean, re.IGNORECASE) or
        re.search(r'^\s*ĐÁP\s+ÁN\s+CHI\s+TIẾT\s*$', text_clean, re.IGNORECASE) or
        re.search(r'^\s*(?:BẢNG\s+)?ĐÁP\s+ÁN\s+CÁC\s+MÃ\s+ĐỀ\s*$', text_clean, re.IGNORECASE)
    )
    if is_title:
        return True
    if len(text_clean) < 50:
        lower_text = text_clean.lower()
        prefixes = ("đáp án", "hướng dẫn chấm", "bảng đáp án", "biểu điểm", "hướng dẫn chấm và biểu điểm", "bảng đáp án", "đáp án và hướng dẫn chấm")
        if any(lower_text.startswith(p) for p in prefixes):
            return True
    if re.search(r'[\-\s_*=]{3,}(?:đáp án|hướng dẫn chấm|biểu điểm)[\-\s_*=]{3,}', text_clean, re.IGNORECASE):
        return True
    return False


def get_answer_for_question(answer_keys: Dict[int, Dict[int, str]], section_idx: int, q_num: int, offsets: Optional[Dict[int, int]] = None) -> Optional[str]:
    """Tìm kiếm đáp án cho câu hỏi trong section_idx, hỗ trợ tìm kiếm fallback nếu đánh số liên tục."""
    if offsets is None:
        offsets = {1: 0, 2: 0, 3: 0}
    if section_idx in answer_keys and q_num in answer_keys[section_idx]:
        return answer_keys[section_idx][q_num]
    abs_q_num = q_num + offsets.get(section_idx, 0)
    for sec_idx in [section_idx, 1, 2, 3]:
        if sec_idx in answer_keys and abs_q_num in answer_keys[sec_idx]:
            return answer_keys[sec_idx][abs_q_num]
    return None


def extract_and_remove_answer_key(doc) -> Dict[int, Dict[int, str]]:
    """Tìm bảng hoặc đoạn văn chứa đáp án ở cuối tài liệu, trích xuất theo phần và xóa bỏ chúng."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    keys = {1: {}, 2: {}, 3: {}}
    body = doc.element.body
    
    # Thu thập tất cả các phần tử (đoạn văn và bảng) theo thứ tự xuất hiện
    elements = []
    for child in body:
        if child.tag.endswith('p'):
            elements.append(("para", Paragraph(child, doc)))
        elif child.tag.endswith('tbl'):
            elements.append(("table", Table(child, doc)))
            
    # Tìm chỉ số tiêu đề bắt đầu vùng đáp án từ dưới lên
    title_idx = -1
    for i in range(len(elements) - 1, -1, -1):
        el_type, el = elements[i]
        if el_type == "para":
            text = el.text.strip()
            if is_answer_title(text):
                title_idx = i
                break
        elif el_type == "table":
            # Kiểm tra xem có ô nào trong bảng chứa tiêu đề đáp án hay không
            found = False
            for row in el.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        cell_txt = p.text.strip()
                        if cell_txt.lower() not in ("đáp án", "câu", "ý"):
                            if is_answer_title(cell_txt):
                                title_idx = i
                                found = True
                                break
                    if found:
                        break
                if found:
                    break
            if found:
                break
                
    elements_to_remove = []
    
    if title_idx != -1:
        current_ans_section_idx = 1
        for idx in range(title_idx, len(elements)):
            el_type, el = elements[idx]
            elements_to_remove.append(el)
            
            if el_type == "para":
                text = el.text.strip()
                sec_match = SECTION_RE.match(text)
                if sec_match:
                    roman_to_int = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, '1': 1, '2': 2, '3': 3, '4': 4}
                    part_str = (sec_match.group(2) or sec_match.group(3)).upper()
                    if part_str in roman_to_int:
                        current_ans_section_idx = roman_to_int[part_str]
                
                # Bỏ qua các đoạn văn giải thích/hướng dẫn giải để tránh khớp nhầm đáp án
                if text.lower().startswith(("giải thích", "lời giải", "hd", "hướng dẫn")):
                    continue
                    
                p_keys = _parse_text_answer_key(text)
                if p_keys:
                    if current_ans_section_idx not in keys:
                        keys[current_ans_section_idx] = {}
                    keys[current_ans_section_idx].update(p_keys)
                    
            elif el_type == "table":
                t_keys = _parse_table_answer_key(el)
                if t_keys:
                    if current_ans_section_idx not in keys:
                        keys[current_ans_section_idx] = {}
                    keys[current_ans_section_idx].update(t_keys)
    else:
        # Fallback quét từ dưới lên
        answer_started = False
        current_ans_section_idx = 1
        for i in range(len(elements) - 1, -1, -1):
            el_type, el = elements[i]
            if el_type == "table":
                t_keys = _parse_table_answer_key(el)
                if t_keys:
                    if current_ans_section_idx not in keys:
                        keys[current_ans_section_idx] = {}
                    keys[current_ans_section_idx].update(t_keys)
                    elements_to_remove.append(el)
                    answer_started = True
                else:
                    if answer_started:
                        break
            elif el_type == "para":
                text = el.text.strip()
                sec_match = SECTION_RE.match(text)
                if sec_match:
                    roman_to_int = {'I': 1, 'II': 2, 'III': 3, 'IV': 4, '1': 1, '2': 2, '3': 3, '4': 4}
                    part_str = (sec_match.group(2) or sec_match.group(3)).upper()
                    if part_str in roman_to_int:
                        current_ans_section_idx = roman_to_int[part_str]
                        
                # Bỏ qua các đoạn văn giải thích/hướng dẫn giải để tránh khớp nhầm đáp án
                if text.lower().startswith(("giải thích", "lời giải", "hd", "hướng dẫn")):
                    if answer_started:
                        elements_to_remove.append(el)
                    continue
                    
                p_keys = _parse_text_answer_key(text)
                if p_keys:
                    if current_ans_section_idx not in keys:
                        keys[current_ans_section_idx] = {}
                    keys[current_ans_section_idx].update(p_keys)
                    elements_to_remove.append(el)
                    answer_started = True
                elif answer_started:
                    is_divider = not text or re.match(r'^[\-\s_.*=]+$', text)
                    is_part_header = len(text) < 100 and (
                        re.match(r'^(phần|part|đáp án|hướng dẫn|câu)', text, re.IGNORECASE) or
                        not re.search(r'\w', text)
                    )
                    is_question = re.match(r'^(câu|question)\s*\d+', text, re.IGNORECASE)
                    is_explanation = text.lower().startswith(("giải thích", "lời giải", "hd", "hướng dẫn"))
                    
                    if is_divider or is_explanation or (is_part_header and not is_question):
                        elements_to_remove.append(el)
                    else:
                        break
                        
    # Xóa các phần tử khỏi XML
    for el in elements_to_remove:
        el_elem = el._element
        parent = el_elem.getparent()
        if parent is not None:
            parent.remove(el_elem)
            
    return keys


def _preprocess_split_runs(p, split_labels, is_lower):
    """Tiền xử lý: Tách một run thành các run nhỏ tại biên giới của các nhãn đáp án mới."""
    suffix_pattern = r'[.:)/\-]'
    p_elem = p._element
    runs = p_elem.findall(qn('w:r'))
    
    for r in runs:
        t_el = r.find(qn('w:t'))
        if t_el is not None and t_el.text:
            text = t_el.text
            for label in split_labels:
                pat = re.compile(rf'(?<![a-zA-Z])({re.escape(label)}\s*{suffix_pattern})')
                m = pat.search(text)
                if m and m.start() > 0:
                    t1 = text[:m.start()]
                    t2 = text[m.start():]
                    t_el.text = t1
                    
                    import copy
                    new_r = copy.deepcopy(r)
                    new_t = new_r.find(qn('w:t'))
                    if new_t is not None:
                        new_t.text = t2
                    
                    r.addnext(new_r)
                    # Đệ quy để tiếp tục tách nếu còn các nhãn khác trong cùng đoạn run
                    _preprocess_split_runs(p, split_labels, is_lower)
                    return


def _merge_adjacent_runs(p) -> bool:
    """Gộp các runs văn bản liên tiếp có cùng định dạng hoặc chỉ chứa text để tránh phân mảnh nhãn."""
    import xml.etree.ElementTree as ET
    p_elem = p._element
    children = list(p_elem)
    i = 0
    modified = False
    while i < len(children) - 1:
        child1 = children[i]
        child2 = children[i+1]
        if child1.tag.endswith('r') and child2.tag.endswith('r'):
            t1 = child1.find(qn('w:t'))
            t2 = child2.find(qn('w:t'))
            
            if t1 is not None and t2 is not None:
                t1_text = t1.text or ""
                t2_text = t2.text or ""
                
                # Cho phép các thẻ định dạng văn bản cơ bản
                has_special1 = any(not (c.tag.endswith('rPr') or c.tag.endswith('t') or c.tag.endswith('tab') or c.tag.endswith('br')) for c in child1)
                has_special2 = any(not (c.tag.endswith('rPr') or c.tag.endswith('t') or c.tag.endswith('tab') or c.tag.endswith('br')) for c in child2)
                
                if has_special1 or has_special2:
                    i += 1
                    continue
                    
                rPr1 = child1.find(qn('w:rPr'))
                rPr2 = child2.find(qn('w:rPr'))
                
                rPr1_str = ET.tostring(rPr1) if rPr1 is not None else b""
                rPr2_str = ET.tostring(rPr2) if rPr2 is not None else b""
                
                same_rpr = (rPr1_str == rPr2_str)
                # Cho phép gộp nếu cùng định dạng hoặc nếu một bên chỉ là dấu chấm
                is_combinable = same_rpr or (t2_text.strip() == ".")
                
                if is_combinable:
                    t1.text = t1_text + t2_text
                    p_elem.remove(child2)
                    children = list(p_elem)
                    modified = True
                    continue
        i += 1
    return modified


def _normalize_paragraph_prefix(p, current_q_num):
    """Chuẩn hóa prefix câu hỏi (ví dụ: Câu 1: -> Câu 1.) mà không phá hủy cấu trúc XML của các run khác."""
    text = p.text
    q_match = QUESTION_RE.match(text)
    if not q_match:
        q_match = PLAIN_NUMBER_RE.match(text)
        
    if not q_match:
        return
        
    prefix_len = q_match.end()
    remainder = text[prefix_len:]
    remainder_clean = remainder.lstrip()
    
    total_replace_len = len(text) - len(remainder_clean)
    
    accumulated_len = 0
    runs_to_modify = []
    
    for run in p.runs:
        run_text = run.text or ""
        runs_to_modify.append((run, run_text))
        accumulated_len += len(run_text)
        if accumulated_len >= total_replace_len:
            break
            
    if runs_to_modify:
        first_run, first_text = runs_to_modify[0]
        if len(first_text) >= total_replace_len:
            first_run.text = f"Câu {current_q_num}. " + first_text[total_replace_len:]
        else:
            first_run.text = f"Câu {current_q_num}. "
            remaining_to_clear = total_replace_len - len(first_text)
            for run, run_text in runs_to_modify[1:]:
                if len(run_text) <= remaining_to_clear:
                    run.text = ""
                    remaining_to_clear -= len(run_text)
                else:
                    run.text = run_text[remaining_to_clear:]
                    break


def _normalize_section_prefix(p, current_section_idx):
    """Chuẩn hóa prefix phần thi (ví dụ: PHẦN I: -> PHẦN 1.) mà không phá hủy cấu trúc XML của các run khác."""
    text = p.text
    sec_match = SECTION_RE.match(text)
    if not sec_match:
        return
        
    prefix_len = sec_match.end()
    remainder = text[prefix_len:]
    remainder_clean = re.sub(r'^[:.\-\s]+', '', remainder)
    
    total_replace_len = len(text) - len(remainder_clean)
    
    accumulated_len = 0
    runs_to_modify = []
    
    for run in p.runs:
        run_text = run.text or ""
        runs_to_modify.append((run, run_text))
        accumulated_len += len(run_text)
        if accumulated_len >= total_replace_len:
            break
            
    if runs_to_modify:
        first_run, first_text = runs_to_modify[0]
        if len(first_text) >= total_replace_len:
            first_run.text = f"PHẦN {current_section_idx}. " + first_text[total_replace_len:]
        else:
            first_run.text = f"PHẦN {current_section_idx}. "
            remaining_to_clear = total_replace_len - len(first_text)
            for run, run_text in runs_to_modify[1:]:
                if len(run_text) <= remaining_to_clear:
                    run.text = ""
                    remaining_to_clear -= len(run_text)
                else:
                    run.text = run_text[remaining_to_clear:]
                    break

def _clean_trailing_whitespace_and_tabs(p):
    """Xóa bỏ các tab, br và khoảng trắng thừa ở cuối đoạn văn (đặc biệt là sau khi tách các phương án)."""
    p_elem = p._element
    while True:
        children = list(p_elem)
        if not children:
            break
        last_c = children[-1]
        if last_c.tag == qn('w:pPr'):
            break
        if last_c.tag == qn('w:r'):
            while True:
                r_children = list(last_c)
                if not r_children:
                    break
                last_r_c = r_children[-1]
                if last_r_c.tag in (qn('w:tab'), qn('w:br'), qn('w:cr'), qn('w:noBreakHyphen')):
                    last_c.remove(last_r_c)
                elif last_r_c.tag == qn('w:t'):
                    if last_r_c.text:
                        last_r_c.text = last_r_c.text.rstrip(' \t\r\n\xa0')
                    if not last_r_c.text:
                        last_c.remove(last_r_c)
                    else:
                        break
                elif last_r_c.tag == qn('w:rPr'):
                    break
                else:
                    break
            meaningful = [c for c in list(last_c) if c.tag != qn('w:rPr')]
            if not meaningful:
                p_elem.remove(last_c)
                continue
            else:
                break
        else:
            break


def _split_inline_answers(doc, p, text: str, is_lower: bool = False) -> bool:
    """Tách các phương án nằm ngang thành các dòng dọc, bảo tồn cấu trúc XML (công thức MathType, định dạng)."""
    pattern = INLINE_TF_RE if is_lower else INLINE_MC_RE
    matches = list(pattern.finditer(text))
    if not matches:
        return False
        
    first_match = matches[0]
    prefix_text = text[:first_match.start()].strip()
    
    if len(prefix_text) > 0:
        split_labels = [m.group(1) for m in matches]
    else:
        split_labels = [m.group(1) for m in matches[1:]]
        
    if not split_labels:
        return False
    
    # Tiền xử lý để tách các run chứa nhiều nhãn
    _preprocess_split_runs(p, split_labels, is_lower)
    
    p_elem = p._element
    parent = p_elem.getparent()
    if parent is None:
        return False
        
    new_paragraphs = []
    
    def create_split_para():
        new_p = doc.add_paragraph()
        pPr = p_elem.find(qn('w:pPr'))
        if pPr is not None:
            import copy
            new_p._element.append(copy.deepcopy(pPr))
        return new_p
        
    current_p = create_split_para()
    new_paragraphs.append(current_p)
    
    label_idx = 0
    
    for child in list(p_elem.iterchildren()):
        if child.tag.endswith('pPr'):
            continue
            
        starts_new_option = False
        if label_idx < len(split_labels):
            next_label = split_labels[label_idx]
            child_text = child.xpath('string(.)')
            suffix_pattern = r'[.:)/\-]'
            match_pattern = rf'^\s*(?:\()?{re.escape(next_label)}\s*{suffix_pattern}'
            if re.match(match_pattern, child_text.strip(), re.IGNORECASE):
                starts_new_option = True
                
        if starts_new_option:
            current_p = create_split_para()
            new_paragraphs.append(current_p)
            label_idx += 1
            
        current_p._element.append(child)
        
    if len(new_paragraphs) <= 1:
        return False
        
    for new_p in new_paragraphs:
        # Xóa các khoảng trắng/tab ở đầu văn bản của run đầu tiên để tránh bị "thò thụt" lệch lề
        for r in new_p._element.findall(qn('w:r')):
            has_text = False
            for child in list(r):
                if child.tag in (qn('w:tab'), qn('w:br')):
                    r.remove(child)
                elif child.tag == qn('w:t'):
                    if child.text and child.text.strip():
                        child.text = child.text.lstrip()
                        has_text = True
                        break
                    elif child.text:
                        child.text = ""
            if has_text:
                break
        # Xóa bỏ các tab và khoảng trắng thừa ở cuối đoạn văn
        _clean_trailing_whitespace_and_tabs(new_p)
        p_elem.addprevious(new_p._element)
        
    parent.remove(p_elem)
    return True


def _add_asterisk_to_label(p, label: str, is_lower: bool = False) -> bool:
    """Thêm dấu hoa thị (*) vào sau nhãn phương án trong XML block."""
    wt_elems = p._element.xpath('.//w:t')
    if not wt_elems:
        return False
        
    # Match ký hiệu nhãn A. hoặc a)
    suffix_pattern = r'[.:)/\-]'
    pat = re.compile(rf'^(\s*)(?:\()?{re.escape(label)}(\s*{suffix_pattern})', re.IGNORECASE)
    
    for wt in wt_elems:
        t = wt.text or ""
        m = pat.match(t)
        if m:
            wt.text = pat.sub(rf'\g<1>{label}*\g<2>', t)
            return True
            
    # Thử quét kiểu thô hơn
    for wt in wt_elems:
        t = wt.text or ""
        if t.strip() == label or t.strip() == f"{label}.":
            wt.text = t.replace(label, f"{label}*")
            return True
            
    return False


def _is_para_marked_as_correct(p) -> bool:
    """Kiểm tra xem phương án có được đánh dấu là đáp án đúng bằng các định dạng đặc biệt hay không (In đậm, Gạch chân, Chữ đỏ, Bôi nền)."""
    for run in p.runs:
        # 1. Kiểm tra in đậm
        if run.bold:
            return True
        try:
            rPr = run._element.get_or_add_rPr()
            if rPr.find(qn('w:b')) is not None:
                return True
                
            # 2. Kiểm tra gạch chân
            if run.underline or rPr.find(qn('w:u')) is not None:
                return True
                
            # 3. Kiểm tra chữ màu đỏ (color = FF0000, red, v.v...)
            color_el = rPr.find(qn('w:color'))
            if color_el is not None:
                val = color_el.get(qn('w:val'))
                # Màu đỏ thuần hoặc các dải màu đỏ
                if val and (val.upper() == 'FF0000' or val.upper() == 'RED'):
                    return True
                    
            # 4. Kiểm tra bôi nền (highlight / shading)
            highlight_el = rPr.find(qn('w:highlight'))
            if highlight_el is not None:
                return True
            shd_el = rPr.find(qn('w:shd'))
            if shd_el is not None:
                fill = shd_el.get(qn('w:fill'))
                if fill and fill.upper() != 'AUTO' and fill.upper() != 'FFFFFF':
                    return True
                    
        except Exception:
            pass
    return False


def normalize_docx(input_path: str, output_path: str, options: dict) -> Tuple[bool, int, str]:
    """Chuẩn hóa file đề thi Word thô."""
    try:
        doc = Document(input_path)

        # Tiền xử lý: Tách toàn bộ các Soft Return (Shift+Enter) thành đoạn văn mới
        split_paragraphs_at_br(doc)

        # 0. Chuyển đổi các bảng chứa các phương án trắc nghiệm thành đoạn văn
        convert_choice_tables_to_paragraphs(doc)

        # 0.5. Dọn dẹp triệt để các đoạn văn rỗng rác (tránh cách dòng thừa)
        clean_empty_paragraphs(doc)

        # 1. Trích xuất và xóa bảng đáp án ở cuối đề
        auto_key = options.get("auto_key", True)
        answer_keys = {}
        if auto_key:
            answer_keys = extract_and_remove_answer_key(doc)
            key_count = sum(len(sec_keys) for sec_keys in answer_keys.values())
            logger.info(f"Đã trích xuất được {key_count} đáp án đúng từ cuối đề.")
        else:
            key_count = 0

        # 1.5. Quét thử một lượt để đếm số lượng câu hỏi trong mỗi phần (để ánh xạ số câu liên tục)
        sec_questions = {1: [], 2: [], 3: []}
        temp_sec_idx = 0
        for p in get_all_paragraphs(doc):
            p_text = p.text.strip()
            if not p_text:
                continue
            sec_match = SECTION_RE.match(p_text)
            if sec_match:
                temp_sec_idx += 1
                continue
        
            if 1 <= temp_sec_idx <= 3:
                q_match = QUESTION_RE.match(p_text)
                if not q_match:
                    q_match = PLAIN_NUMBER_RE.match(p_text)
                if q_match:
                    try:
                        if q_match.re == QUESTION_RE:
                            q_num = int(q_match.group(3))
                        else:
                            q_num = int(q_match.group(1))
                        sec_questions[temp_sec_idx].append(q_num)
                    except Exception:
                        pass
                    
        is_already_continuous = True
        if sec_questions[1] and sec_questions[2]:
            if min(sec_questions[2]) <= max(sec_questions[1]):
                is_already_continuous = False
        if is_already_continuous and sec_questions[2] and sec_questions[3]:
            if min(sec_questions[3]) <= max(sec_questions[2]):
                is_already_continuous = False
            
        offsets = {1: 0, 2: 0, 3: 0}
        if not is_already_continuous:
            offsets[2] = len(sec_questions[1])
            offsets[3] = len(sec_questions[1]) + len(sec_questions[2])
        
        logger.info(f"Đã tính toán offsets câu hỏi: {offsets}")

        # 2. Duyệt qua tất cả các đoạn văn để chuẩn hóa
        current_section_idx = 0
        current_q_num = 0
        p_idx = 0
    
        # Tạo danh sách các đoạn văn để tránh lỗi kích thước mảng thay đổi khi insert/remove
        paragraphs = get_all_paragraphs(doc)
        # Tập hợp id() của các phần tử XML đã được split, để không split lại sau khi reload paragraphs
        split_done_ids = set()
    
        while p_idx < len(paragraphs):
            p = paragraphs[p_idx]
            _merge_adjacent_runs(p)
            text = p.text.strip()
        
            if not text:
                p_idx += 1
                continue
            
            # A. Kiểm tra tiêu đề phần
            sec_match = SECTION_RE.match(text)
            if sec_match:
                # Chuẩn hóa tiêu đề phần thành "PHẦN 1...", "PHẦN 2..."
                # Tự động tăng phần thi
                current_section_idx += 1
                _normalize_section_prefix(p, current_section_idx)
                if p.runs:
                    p.runs[0].bold = True
                p_idx += 1
                continue
            
            # B. Kiểm tra câu hỏi
            q_match = QUESTION_RE.match(text)
            if not q_match:
                # Thử định dạng số thuần túy như "1." ở đầu đoạn văn
                q_match = PLAIN_NUMBER_RE.match(text)
            
            if q_match:
                if q_match.re == QUESTION_RE:
                    current_q_num = int(q_match.group(3))
                else:
                    current_q_num = int(q_match.group(1))
                # Chuẩn hóa về Câu X.
                _normalize_paragraph_prefix(p, current_q_num)
                # KHÔNG continue ở đây để nếu có phương án A. nằm cùng dòng với câu hỏi, nó sẽ được xử lý ở bước E.
            
            # E. Phát hiện và xử lý đáp án nằm ngang (chạy trước khi kiểm tra từng phương án đơn lẻ)
            elem_id = id(p._element)
            if current_q_num > 0 and elem_id not in split_done_ids:
                if _split_inline_answers(doc, p, text, is_lower=(current_section_idx == 2)):
                    split_done_ids.add(elem_id)
                    # Reload lại danh sách paragraphs do cấu trúc docx đã bị thay đổi (xóa và thêm p)
                    paragraphs = get_all_paragraphs(doc)
                    # KHÔNG tăng p_idx: xử lý lại đoạn đầu tiên vừa tách để đánh dấu đáp án đúng
                    # split_done_ids đảm bảo đoạn đó không bị split lại vô hạn
                    continue

            # C. Kiểm tra phương án trắc nghiệm (Nhận dạng nhãn A. B. C. D.)
            m_upper = re.match(r'^(?:\()?([A-H])\s*[.:)/\-]\s*', text)
        
            if m_upper and current_q_num > 0:
                label = m_upper.group(1).upper()
                is_correct = False
            
                # Check theo đáp án đúng trích xuất
                correct_val = get_answer_for_question(answer_keys, current_section_idx, current_q_num, offsets)
                if correct_val and label == correct_val.upper():
                    is_correct = True
                    
                # Check theo in đậm/gạch chân/màu sắc (chỉ dùng khi không có đáp án từ key)
                if not correct_val and options.get("bold_as_correct", False) and _is_para_marked_as_correct(p):
                    is_correct = True
                
                if is_correct:
                    # Thêm dấu hoa thị (*)
                    _add_asterisk_to_label(p, label, is_lower=False)
            
                p_idx += 1
                continue
            
            # D. Kiểm tra mệnh đề Đúng/Sai Phần 2
            m_lower = re.match(r'^([a-h])\s*[\).]\s*', text)
            if m_lower and current_q_num > 0:
                label = m_lower.group(1).lower()
                is_correct = False
            
                # Check theo đáp án Đúng/Sai trích xuất (dạng Đ,S,Đ,S)
                correct_val = get_answer_for_question(answer_keys, current_section_idx, current_q_num, offsets)
                if correct_val:
                    vals = [x.strip() for x in correct_val.split(',')]
                    idx_val = ord(label) - ord('a')
                    if idx_val < len(vals) and vals[idx_val] in ['Đ', 'D']:
                        is_correct = True
                    
                # Check theo in đậm (chỉ dùng khi không có đáp án từ key)
                if not correct_val and options.get("bold_as_correct", False) and _is_para_marked_as_correct(p):
                    is_correct = True
                
                if is_correct:
                    _add_asterisk_to_label(p, label, is_lower=True)
                
                p_idx += 1
                continue
                
            p_idx += 1
        
        # 3. Chèn đáp án cho Phần 3 (Trả lời ngắn) nếu chưa có
        # Duyệt lại toàn bộ để chèn Đáp án: [Kết quả] cho các câu Phần 3
        if current_section_idx >= 3 or key_count > 0:
            p_idx = 0
            paragraphs = get_all_paragraphs(doc)
            while p_idx < len(paragraphs):
                p = paragraphs[p_idx]
                text = p.text.strip()
                
                q_match = QUESTION_RE.match(text)
                if q_match:
                    q_num = int(q_match.group(3))
                
                    # Xác định phần thi của câu hỏi này
                    sec_num = 0
                    for j in range(p_idx, -1, -1):
                        sec_match = SECTION_RE.match(paragraphs[j].text.strip())
                        if sec_match:
                            sec_num = int(sec_match.group(2))
                            break
                        
                    correct_val = get_answer_for_question(answer_keys, sec_num, q_num, offsets)
                    if correct_val and (sec_num == 3 or is_short_answer_value(correct_val)):
                        # Kiểm tra xem đoạn kế tiếp có phải là dòng "Đáp án: ..." không
                        has_answer_line = False
                        if p_idx + 1 < len(paragraphs):
                            next_text = paragraphs[p_idx + 1].text.strip()
                            if next_text.startswith("Đáp án"):
                                has_answer_line = True
                            
                        if not has_answer_line:
                            # Chèn đoạn văn mới "Đáp án: [Kết quả]" ngay sau câu hỏi
                            p_elem = p._element
                            parent = p_elem.getparent()
                            if parent is not None:
                                new_p = doc.add_paragraph()
                                new_p.text = f"Đáp án: {correct_val}"
                                p_elem.addnext(new_p._element)
                                # Cập nhật danh sách paragraphs
                                paragraphs = get_all_paragraphs(doc)
                                p_idx += 1
                p_idx += 1

        # 4. Tùy chọn dồn dòng phương án
        choice_layout = options.get("choice_layout", "split")
        choice_indent = options.get("choice_indent", True)
        indent_dxa = options.get("indent_dxa", 567) if choice_indent else 0
        compacted_msg = ""
        if choice_layout and (choice_layout != "split" or indent_dxa > 0):
            from choice_compactor import compact_document_choices
            total_g, compacted_count = compact_document_choices(doc, choice_layout, indent_dxa)
            if compacted_count > 0:
                compacted_msg = f" Đã dồn dòng {compacted_count} câu hỏi."

        # Lưu tài liệu kết quả
        doc.save(output_path)
        return True, key_count, f"Chuẩn hóa thành công! Đã quét được {key_count} đáp án.{compacted_msg}"
        
    except Exception as e:
        logger.exception("Lỗi khi chuẩn hóa đề thi: %s", e)
        return False, 0, str(e)
