"""
Decree 30/2020/NĐ-CP Document Normalizer for ExamScrambler Ultimate.
Provides automated compliance formatting for general administrative documents.
"""
import os
import re
import logging
from typing import Dict, Any
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

logger = logging.getLogger(__name__)

# Regex patterns for detecting headings and sections
HEADING_RE = re.compile(
    r'^\s*('
    r'Chương\s+[IVXivx\d]+'
    r'|Mục\s+\d+'
    r'|Điều\s+\d+'
    r'|[A-ZĐ\d]+\.\s+'   # I., II., 1., 2.
    r'|[a-zđ]\)\s+'       # a), b)
    r'|Part\s+[A-Z\d]'
    r')',
    re.IGNORECASE
)

MAIN_HEADING_RE = re.compile(
    r'^\s*('
    r'Chương\s+[IVXivx\d]+'
    r'|Mục\s+\d+'
    r'|Điều\s+\d+'
    r'|[A-ZĐ\d]+\.\s+'
    r')',
    re.IGNORECASE
)

SUB_HEADING_RE = re.compile(r'^\s*([a-zđ]\)\s+)', re.IGNORECASE)


def set_run_font(run, font_name: str = 'Times New Roman', font_size_pt: float = None):
    """Sets Times New Roman font for ASCII, EastAsia, hAnsi, and CS text in run XML to prevent Word fallbacks."""
    run.font.name = font_name
    if font_size_pt is not None:
        run.font.size = Pt(font_size_pt)
        
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rPr.append(rFonts)


def add_page_number_to_header(section):
    """Adds a centered, standard page number field in the header (Arabic numerals, starting from page 2)."""
    # Enable separate header for first page so page 1 is not numbered
    section.different_first_page_header_footer = True
    
    header = section.header
    if not header.paragraphs:
        p = header.add_paragraph()
    else:
        p = header.paragraphs[0]
        p.text = ""  # Clear previous content

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    # Simple simpleField instruction for page number in Word
    run = p.add_run()
    set_run_font(run, font_name='Times New Roman', font_size_pt=13)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)


def normalize_document_nd30(input_path: str, output_path: str, options: Dict[str, Any]) -> str:
    """
    Reads a docx document, formats it according to Decree 30 specifications, and saves it.
    """
    logger.info(f"Bắt đầu chuẩn hóa file {input_path} theo NĐ 30")
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Không tìm thấy file nguồn: {input_path}")
        
    doc = Document(input_path)
    
    # 1. Page Setup & Margins
    if options.get("margin", True):
        logger.info("Thiết lập kích thước trang A4 và căn lề theo NĐ 30")
        for section in doc.sections:
            section.page_width = Cm(21.0)
            section.page_height = Cm(29.7)
            section.top_margin = Cm(2.0)
            section.bottom_margin = Cm(2.0)
            section.left_margin = Cm(3.0)
            section.right_margin = Cm(1.5)

    # 2. Page Numbering (Top Center, skipping page 1)
    if options.get("page_num", True):
        logger.info("Thêm đánh số trang tự động ở đầu trang")
        for section in doc.sections:
            add_page_number_to_header(section)

    # 3. Format Body Paragraphs
    if options.get("font", True) or options.get("paragraph", True) or options.get("headings", True):
        logger.info("Căn chỉnh phông chữ, đoạn văn và tiêu đề đề mục")
        
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            # Detect National motto / Slogans
            text_lower = text.lower()
            is_slogan1 = "cộng hòa xã hội chủ nghĩa việt nam" in text_lower
            is_slogan2 = "độc lập - tự do - hạnh phúc" in text_lower
            
            if is_slogan1 or is_slogan2:
                # Align Center, Bold, sizes 13/14
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = None
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                size = 13 if is_slogan1 else 14
                for run in p.runs:
                    set_run_font(run, font_name='Times New Roman', font_size_pt=size)
                    run.font.bold = True
                    run.font.italic = False
                continue

            # Detect document main title (short centered lines in all uppercase, e.g., QUYẾT ĐỊNH, KẾ HOẠCH)
            is_main_title = text.isupper() and len(text) < 100
            
            # Detect standard headings (Chương, Mục, Điều, 1., a)...)
            is_heading = HEADING_RE.match(text)
            
            if is_main_title or is_heading:
                # Apply Heading rules
                if options.get("headings", True):
                    p.paragraph_format.first_line_indent = None
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.line_spacing = 1.15
                    
                    if is_main_title:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in p.runs:
                            set_run_font(run, font_name='Times New Roman', font_size_pt=14)
                            run.font.bold = True
                            run.font.italic = False
                    elif MAIN_HEADING_RE.match(text):
                        # Keep alignment or justify
                        for run in p.runs:
                            set_run_font(run, font_name='Times New Roman', font_size_pt=14)
                            run.font.bold = True
                            run.font.italic = False
                    elif SUB_HEADING_RE.match(text):
                        # Italicize lower level list item (a, b)
                        for run in p.runs:
                            set_run_font(run, font_name='Times New Roman', font_size_pt=13)
                            run.font.bold = False
                            run.font.italic = True
                    else:
                        # Fallback for general heading
                        for run in p.runs:
                            set_run_font(run, font_name='Times New Roman', font_size_pt=13)
                            run.font.bold = True
                else:
                    # Just standardize font if heading formatting is disabled
                    if options.get("font", True):
                        for run in p.runs:
                            set_run_font(run, font_name='Times New Roman')
                continue

            # Format general body paragraph
            if options.get("paragraph", True):
                if p.alignment != WD_ALIGN_PARAGRAPH.CENTER and p.alignment != WD_ALIGN_PARAGRAPH.RIGHT:
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    # First line indent standard: 1.0 cm to 1.27 cm
                    p.paragraph_format.first_line_indent = Cm(1.0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.2

            if options.get("font", True):
                for run in p.runs:
                    # Retain original size if it was 13 or 14, otherwise default to 14
                    orig_size = run.font.size
                    target_size = 14.0
                    if orig_size is not None:
                        size_pt = orig_size.pt
                        if size_pt in [13.0, 14.0]:
                            target_size = size_pt
                    set_run_font(run, font_name='Times New Roman', font_size_pt=target_size)

    # 4. Format Tables
    if options.get("font", True):
        logger.info("Căn chỉnh phông chữ trong các bảng biểu")
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        # In tables, typically no first line indent is used
                        p.paragraph_format.first_line_indent = None
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.line_spacing = 1.15
                        
                        for run in p.runs:
                            # Typically tables use slightly smaller text (size 12 or 13)
                            orig_size = run.font.size
                            target_size = 12.0
                            if orig_size is not None:
                                size_pt = orig_size.pt
                                if size_pt in [12.0, 13.0, 14.0]:
                                    target_size = size_pt
                            set_run_font(run, font_name='Times New Roman', font_size_pt=target_size)

    # Save to output file
    doc.save(output_path)
    logger.info(f"Đã hoàn thành chuẩn hóa văn bản NĐ 30, lưu tại {output_path}")
    return output_path
